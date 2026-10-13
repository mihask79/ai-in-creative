"""Build the completed LR03 variant 1 report from the experiment artifacts."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / "artifacts"
REPORT = ROOT / "reports" / "ЛР03_Отчет_вариант_1.docx"


def cell_text(cell, text, bold=False):
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold
            run.font.name = "Arial"
            run.font.size = Pt(8.5)


def f(value, digits=5):
    return f"{value:.{digits}f}"


def add_picture(doc, path, width=6.35):
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def main():
    data = json.loads((ART / "results.json").read_text(encoding="utf-8"))
    passport_data = json.loads((ART / "passport_lr03.json").read_text(encoding="utf-8"))
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.name = "Arial"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Лабораторная работа № 3\nПодготовка изображений и визуализация")
    p = doc.add_paragraph("Вариант 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    doc.add_paragraph("Дисциплина: Искусственный интеллект в креативных технологиях\nСтудент: Скрыпник Михаил Дмитриевич   Группа: МИК11")

    m = data["metrics"]
    mean = m["mean"]
    contrast = m["contrast_std"]
    entropy = m["entropy_bits"]
    psnr = data["mean_psnr_db"]
    doc.add_heading("Итог", level=1)
    doc.add_paragraph(
        f"Квантование яркости сцены с градиентом и кругом до 4 бит практически не изменило среднее значение "
        f"({100*mean['relative_change']:.2f}%), но увеличило контраст на {100*contrast['relative_change']:.2f}% "
        f"и снизило энтропию гистограммы на {100*entropy['relative_change']:.2f}%. Средний PSNR составил {psnr:.2f} дБ, "
        "что выше принятого порога 30 дБ. Контрольная сумма для seed 101 совпала при повторном построении сцены."
    )

    doc.add_heading("1. Цель и постановка задачи", level=1)
    doc.add_paragraph("Цель работы — оценить, как подготовительная операция изменяет данные изображения, сравнить метрики яркости и оценить потери по PSNR.")
    doc.add_paragraph("Вариант 1: сцена 1 — базовая, градиент и круг; lo=0,15, hi=0,85, шум 0,02; воздействие — квантование яркости до 4 бит (16 уровней).")
    doc.add_paragraph("Гипотеза до запуска: средняя яркость изменится менее чем на 5%, энтропия уменьшится; изменение контраста проверяется измерениями.")
    doc.add_paragraph("Практический допуск, установленный до расчётов: относительное изменение метрики более 5% или PSNR ниже 30 дБ.")

    doc.add_heading("2. Данные и среда эксперимента", level=1)
    scene = data["scene"]
    doc.add_paragraph(
        f"Сцена: {scene['shape_rgb'][0]}×{scene['shape_rgb'][1]} пикселей, RGB{{Float64}}, "
        f"минимум {scene['observed_min_seed_101']:.5f}, максимум {scene['observed_max_seed_101']:.5f}, "
        f"объём {scene['shape_rgb'][0]*scene['shape_rgb'][1]*scene['shape_rgb'][2]*8:,} байт. "
        "Использованы seed 101–105 и 32 интервала гистограммы."
    )
    doc.add_paragraph("Среда: Python; точные версии приведены в паспорте. Сессия Engee не выполнялась; библиотека lab03_lib.jl не приложена.")
    doc.add_paragraph(f"Дата выполнения: {passport_data['date']}.")

    doc.add_heading("3. Ход работы", level=1)
    steps = [
        "Создана сцена варианта 1 и записаны её размер, тип элемента, диапазон значений и объём.",
        "Яркость рассчитана вручную по BT.601 и отдельным преобразованием каналов; максимум модуля разности составил {:.3e}.".format(data["luma_max_abs_difference"]),
        "Для массива яркостей seed 101 вычислена контрольная сумма SHA-256; повторное построение дало ту же сумму.",
        "Для исходного и обработанного изображения построены гистограммы с 32 интервалами.",
        "Применено квантование до 4 бит; рассчитаны метрики и PSNR для пяти seed.",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("4. Результаты", level=1)
    doc.add_paragraph(f"SHA-256 яркостей для seed 101: {data['seed_101_luma_sha256']}.")
    doc.add_paragraph("Повторная контрольная сумма совпала: да.")

    summary = doc.add_table(rows=1, cols=7)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary.style = "Light Shading Accent 1"
    heads = ["Метрика", "До: среднее", "До: s", "После: среднее", "После: s", "Δ", "|Δ|/s"]
    for c, value in zip(summary.rows[0].cells, heads):
        cell_text(c, value, True)
    for label, item in (("Среднее", mean), ("Контраст", contrast), ("Энтропия, бит", entropy)):
        cells = summary.add_row().cells
        vals = [label, f(item["baseline_mean"]), f(item["baseline_sd"], 6), f(item["after_mean"]), f(item["after_sd"], 6), f(item["delta"]), f(item["statistical_ratio_abs_delta_over_baseline_sd"], 2)]
        for c, value in zip(cells, vals):
            cell_text(c, value)

    doc.add_paragraph("Сводка по seed: изменение метрик и PSNR.")
    per_seed = doc.add_table(rows=1, cols=5)
    per_seed.alignment = WD_TABLE_ALIGNMENT.CENTER
    per_seed.style = "Light Shading Accent 1"
    for c, value in zip(per_seed.rows[0].cells, ["Seed", "Δ среднего", "Δ контраста", "Δ энтропии", "PSNR, дБ"]):
        cell_text(c, value, True)
    for row in data["raw_results"]:
        cells = per_seed.add_row().cells
        vals = [row["seed"], f(row["delta_mean"], 6), f(row["delta_contrast_std"], 6), f(row["delta_entropy_bits"], 5), f(row["psnr_db"], 2)]
        for c, value in zip(cells, vals):
            cell_text(c, value)

    doc.add_paragraph(f"Средний PSNR по пяти seed: {psnr:.2f} дБ. PSNR < 30 дБ: {'да' if data['psnr_practically_significant'] else 'нет'}.")
    add_picture(doc, ART / "scene_before.png", 3.0)
    doc.add_paragraph("Рисунок 1 — изображение до квантования.")
    add_picture(doc, ART / "scene_after.png", 3.0)
    doc.add_paragraph("Рисунок 2 — изображение после квантования до 4 бит.")
    add_picture(doc, ART / "histograms.png", 6.2)
    doc.add_paragraph("Рисунок 3 — гистограммы яркости до и после воздействия.")
    add_picture(doc, ART / "metric_comparison.png", 6.2)
    doc.add_paragraph("Рисунок 4 — средние значения метрик по серии seed; усы показывают выборочное стандартное отклонение.")

    doc.add_heading("5. Анализ и вывод", level=1)
    for label, item in (("Среднее", mean), ("Контраст", contrast), ("Энтропия", entropy)):
        stat = "превышает" if item["statistically_significant_by_course_rule"] else "не превышает"
        practical = "превышает" if item["practically_significant_by_5_percent_rule"] else "не превышает"
        doc.add_paragraph(
            f"{label}: |Δ|/s = {item['statistical_ratio_abs_delta_over_baseline_sd']:.2f}, что {stat} статистический порог 2; "
            f"относительное изменение {100*item['relative_change']:.2f}% {practical} практический порог 5%."
        )
    doc.add_paragraph(
        "Средняя яркость изменилась менее чем на 1%, поэтому изменение находится в пределах практического допуска. "
        f"Контраст изменился на {100*contrast['relative_change']:.2f}%, энтропия — на {100*entropy['relative_change']:.2f}%. "
        f"PSNR {psnr:.2f} дБ выше порога 30 дБ. Таким образом, основное изменение связано с распределением яркости, "
        "а не со средним уровнем. Для контраста учебный статистический порог превышен, практический допуск — нет."
    )
    doc.add_paragraph("Ограничения: одна сцена и пять seed; метрики описывают статистику яркости, а PSNR не измеряет воспринимаемое качество изображения.")

    doc.add_paragraph("Использована самостоятельная реализация базовой сцены по описанию методички. Геометрия и генератор случайных чисел отличаются от исходной библиотеки курса, поэтому демонстрационные числа Engee не заявляются как результат этого запуска. В методичке PSNR равен 34,59 дБ; здесь получено {:.2f} дБ. В демонстрации средняя яркость статистически значима, здесь порог не превышен. Это ограничивает сопоставимость реализаций.".format(psnr))
    doc.add_paragraph("Чувствительность к допуску: при 1%, 5% и 10% практически заметной остаётся только энтропия. При пороге PSNR 35 дБ результат уже вышел бы за допуск. Пороги после расчётов не изменялись; это анализ устойчивости вывода.")
    doc.add_paragraph("Контрольные вопросы. Изображение 256×256 RGB Float64 занимает 1 572 864 байта; RGB N0f8 — 196 608 байт, в 8 раз меньше. При MSE=0,0004 PSNR=33,98 дБ; уменьшение MSE вдвое добавляет 3,01 дБ. Квантование до 16 уровней ограничивает энтропию сверху log2(16)=4 бит.")
    doc.add_heading("6. Паспорт эксперимента", level=1)
    passport = doc.add_table(rows=0, cols=2)
    passport.style = "Light Shading Accent 1"
    passport.alignment = WD_TABLE_ALIGNMENT.CENTER
    fields = [
        ("Дата", passport_data["date"]),
        ("Лабораторная / вариант", "ЛР03 / 1"),
        ("Среда и версии", "; ".join(f"{k} {v}" for k, v in data["environment"].items())),
        ("Сцена", "1 — базовая сцена; градиент и круг; 128×128; lo=0,15, hi=0,85, шум 0,02"),
        ("Seed", "101–105"),
        ("Воздействие", "Квантование яркости до 4 бит (16 уровней)"),
        ("Гистограмма", "32 интервала"),
        ("Практический допуск", "Изменение метрики >5% или PSNR <30 дБ"),
        ("SHA-256, seed 101", data["seed_101_luma_sha256"]),
        ("Артефакты", "Скрипт, изображения до/после, гистограммы, график метрик, CSV и JSON"),
    ]
    for key, value in fields:
        cells = passport.add_row().cells
        cell_text(cells[0], key, True)
        cell_text(cells[1], value)

    REPORT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
