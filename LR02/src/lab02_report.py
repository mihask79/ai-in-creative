"""Build the variant 1 report from the current audit and model card."""
from pathlib import Path
import csv
import json
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]

def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.8)
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.1
    for name in ['Title', 'Heading 1', 'Heading 2']:
        doc.styles[name].font.name = 'Arial'
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles[name].font.size = Pt({'Title':20, 'Heading 1':14, 'Heading 2':11}[name])
        doc.styles[name].paragraph_format.space_before = Pt(10)
        doc.styles[name].paragraph_format.space_after = Pt(5)
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.core_properties.title = 'Лабораторная работа 2 Карточка модели лицензии и риски Вариант 1'
    doc.core_properties.author = 'Скрыпник Михаил Дмитриевич'
    for node in doc.styles.element.xpath('.//w:pBdr'):
        node.getparent().remove(node)
    doc.add_paragraph('Карточка модели лицензии и риски', style='Title')
    doc.add_paragraph('Лабораторная работа 2 · Вариант 1\nДисциплина: Искусственный интеллект в креативных технологиях\nСтудент: Скрыпник Михаил Дмитриевич · Группа: МИК11\nДата исходного аудита: 13.10.2026\nДата актуализации отчёта: 27.10.2026')
    doc.add_heading('1 Цель и сценарий', level=1)
    doc.add_paragraph('Проверить пригодность набора для коммерческой афиши фестиваля: печать, продажа билетов, публикация и создание производных. ShareAlike не допускается. Аудит выявил четыре ошибки и три предупреждения. Для четырёх объектов принято решение об исключении до отдельного разрешения, один ожидает подтверждения даты; семь оставшихся записей условно допустимы в рамках учебных правил. Исходный манифест сохранён полностью для воспроизведения проверки.')
    doc.add_paragraph('Исследовались 12 синтетических записей манифеста, без реальных файлов изображений. Карточка модели относится к процедурному генератору облачного фона из ЛР 1, вариант 1. Ссылки на входные данные, код, результаты и изображения генератора приведены в разделе 9.')
    doc.add_paragraph('До запуска ожидались конфликты NC у item01, item04, item08, item10, конфликты SA у item01 и item04 и отсутствующая дата item09. Ожидания записаны в reports/lab_journal.txt. Критерий успеха — решение по каждой находке, заполненные карточки и проверка рисков PASS.')
    doc.add_heading('2 Методика и воспроизведение', level=1)
    doc.add_paragraph('Использован Python 3.12.14 и стандартная библиотека для аудита. Для сборки DOCX применяется python-docx. Манифест синтетический: реальные произведения не загружались. Учебные правила из методички проверяют поля записей, а не подлинность лицензии. Команды выполняются из папки LR02.')
    for command in [
        'python src/lab02_make_variant.py --variant 1 --out artifacts',
        'python src/lab02_audit.py --manifest artifacts/manifest.csv --policy artifacts/policy.json --out reports',
        'python src/lab02_model_card.py --passport artifacts/passport_lr01.json --out reports/model_card.md',
        'python src/lab02_risks.py --csv artifacts/risk_register.csv --out reports/risk_ranking.md',
    ]:
        p=doc.add_paragraph(command); p.runs[0].font.size=Pt(9)
    doc.add_heading('3 Находки и решения', level=1).paragraph_format.page_break_before = True
    with (ROOT/'reports/audit_findings.csv').open(encoding='utf-8-sig', newline='') as f:
        findings=list(csv.DictReader(f))
    for item in dict.fromkeys(r['item_id'] for r in findings):
        subset=[r for r in findings if r['item_id']==item]
        doc.add_heading(item+' '+' '.join(r['rule'] for r in subset), level=2)
        doc.add_paragraph('; '.join(r['description'] for r in subset)+'.')
        doc.add_paragraph(subset[0]['decision'])
    doc.add_paragraph('Решения принимаются по объектам: исключение item01 и item04 снимает обе находки по каждому. Сценарий не допускает коммерческое использование NC и передачу производных на условиях SA. Отдельное разрешение правообладателя должно охватывать все конфликтующие условия.')
    doc.add_paragraph('Итого: четыре ошибки R3, два предупреждения R5 и одно предупреждение R8 по пяти объектам. Условно допустимы item02, item03, item05, item06, item07, item11 и item12. После исключения четырёх NC-объектов остаётся восемь записей, но item09 находится в карантине. Подтверждение его даты и происхождения может увеличить допустимый состав с семи до восьми записей.')
    doc.add_heading('4 Паспорт набора данных', level=1).paragraph_format.page_break_before = True
    add_markdown(doc, ROOT/'reports/data_card.md')
    doc.add_heading('5 Карточка модели', level=1).paragraph_format.page_break_before = True
    add_markdown(doc, ROOT/'reports/model_card.md')
    doc.add_heading('6 Реестр рисков', level=1).paragraph_format.page_break_before = True
    with (ROOT/'artifacts/risk_register.csv').open(encoding='utf-8-sig',newline='') as f:
        risks=list(csv.DictReader(f))
    for risk in sorted(risks,key=lambda r:int(r['likelihood'])*int(r['impact']),reverse=True):
        score=int(risk['likelihood'])*int(risk['impact'])
        doc.add_heading(f"{risk['risk_id']} {risk['category']}",level=2)
        doc.add_paragraph(f"Вероятность {risk['likelihood']}, ущерб {risk['impact']}, приоритет {score}. {risk['description']}.")
        doc.add_paragraph(f"Мера: {risk['mitigation']}. Ответственный: {risk['owner']}. Триггер: {risk['trigger']}.")
    doc.add_paragraph('Проверка: RESULT: PASS. Шесть рисков в шести категориях; обязательные Intellectual Property и Information Integrity присутствуют. Оценка P×I задаёт учебный приоритет и не является вероятностью ущерба.')
    doc.add_heading('7 Диагностика ошибок и границы вывода', level=1).paragraph_format.page_break_before = True
    doc.add_paragraph('В отдельной проверке намеренно задана категория Copyright, оценка 7 и пустой ответственный. Валидатор отклоняет такой реестр. Исправление: Intellectual Property, целочисленная оценка 1–5 и заполненный ответственный; рабочий реестр проходит проверку. Протокол записан в reports/validation.txt.')
    doc.add_paragraph('Инженерное решение — исключить конфликтующие записи, хранить атрибуцию и вести карантин. Юридический вопрос — достаточно ли конкретного разрешения для печати, публикации и производных и как квалифицируется переработка; он требует проверки реальных документов. Этическое решение — не вводить зрителя в заблуждение об авторах и программе фестиваля. Организационное решение — назначить куратора и редактора, запретить выпуск до закрытия находок.')
    doc.add_paragraph('Контрольные вопросы. При переходе к некоммерческому сценарию R3 может исчезнуть, но условия SA и требование подтверждения происхождения сохраняются. Отсутствие ошибки в автоматическом аудите не подтверждает права: скрипт видит только записанные поля. По объекту с NC-SA нужно либо исключение, либо разрешение, снимающее оба конфликта.')
    doc.add_heading('8 Вывод и материалы',level=1)
    doc.add_paragraph('Для текущей политики исходный набор не готов к использованию. Решения по всем семи находкам записаны; карточка генератора использует фактические результаты варианта 1 из LR01-01. Ограничения: синтетические записи, упрощённые правила и неизвестные лицензия кода и ответственный за его сопровождение. Полный комплект: assignment/, src/, artifacts/, reports/ и README.md.')
    add_artifacts(doc)
    path=ROOT/'reports/ЛР02_Отчет_вариант_1.docx'
    doc.save(path)
    print(path)

def add_markdown(doc, path):
    headings = {
        'Model Details':'Сведения о генераторе', 'Intended Use':'Назначение',
        'Factors':'Факторы', 'Metrics':'Метрики', 'Evaluation Data':'Данные для оценки',
        'Training Data':'Данные для обучения', 'Quantitative Analyses':'Количественные результаты',
        'Ethical Considerations':'Этические аспекты', 'Caveats and Recommendations':'Ограничения и рекомендации',
    }
    current = ''
    for line in path.read_text(encoding='utf-8').splitlines():
        if line.startswith('# '): continue
        if line.startswith('## '):
            current = line[3:]
            doc.add_heading(headings.get(current, current.replace('.', '')),level=2)
        elif line.strip():
            if current == 'Model Details':
                line = 'Учебный процедурный генератор версии 1.0 из LR01-01/src/lab01_generator.py. Нейросеть и обученные веса отсутствуют. Среда исходного эксперимента: CPython 3.12.14, Windows 11, zlib 1.3.2.'
            if current == 'Quantitative Analyses':
                add_metrics(doc)
            else:
                doc.add_paragraph(line)


def add_metrics(doc):
    data = json.loads((ROOT.parent/'LR01-01/artifacts/results.json').read_text(encoding='utf-8'))
    assert data['variant'] == 1
    doc.add_paragraph('Данные из LR01-01/artifacts/results.json. По пять seed (101–105) для базы и фактора; s — выборочное стандартное отклонение базовых измерений. Учебный критерий существенного эффекта: |Δ|/s > 2.')
    t = doc.add_table(rows=1, cols=5)
    t.autofit = False
    widths = [4.6, 3.0, 3.0, 3.0, 3.3]
    for c,w in zip(t.columns,widths): c.width = Cm(w)
    for c,v in zip(t.rows[0].cells,['Метрика','База','s базы','Фактор','|Δ|/s']): c.text=v
    header = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(header)
    for key,label in [('mean','Средняя яркость'),('contrast','Контраст'),('edge_energy','Граничная энергия'),('entropy_bits','Энтропия бит')]:
        m=data['metrics'][key]
        vals=[label]+[f'{m[k]:.6f}'.replace('.',',') for k in ['baseline_mean','baseline_sample_sd','factor_mean']]+[f"{m['abs_delta_over_baseline_sd']:.3f}".replace('.',',')]
        for c,v in zip(t.add_row().cells,vals): c.text=v
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: e.set(qn('w:'+k),v)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c,w in zip(row.cells,widths):
            c.width=Cm(w)
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'DDE8F1' if i==0 else 'FFFFFF'); pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'75'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.line_spacing=1.0
                for r in p.runs: r.font.size=Pt(9); r.bold=(i==0)
    doc.add_paragraph('Ни одна метрика не превысила порог 2. Вывод ограничен данным генератором, пятью seed и изменением размера 96→144.')


def add_link(doc, relative, description):
    target=(ROOT/'reports'/relative).resolve()
    if not target.is_file(): raise FileNotFoundError(target)
    p=doc.add_paragraph()
    h=OxmlElement('w:hyperlink')
    h.set(qn('r:id'), p.part.relate_to(relative, RT.HYPERLINK, is_external=True))
    r=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'175A8A'); props.append(color)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); props.append(u)
    size=OxmlElement('w:sz'); size.set(qn('w:val'),'19'); props.append(size)
    r.append(props); text=OxmlElement('w:t'); text.text=relative; r.append(text); h.append(r); p._p.append(h)
    p.add_run(' — '+description).font.size=Pt(9.5)
    p.paragraph_format.space_after=Pt(5)


def add_artifacts(doc):
    doc.add_heading('9 Расположение исследованных материалов',level=1).paragraph_format.page_break_before=True
    doc.add_paragraph('Рабочая папка: C:\\DONSTU\\ai-in-creative-industries\\LR02. Отчёт находится в reports/ЛР02_Отчет_вариант_1.docx. Все ссылки ниже заданы относительно папки отчёта. Для передачи подготовлен архив LR02_variant_1.zip: его следует распаковать с сохранением структуры и открыть отчёт в LR02/reports/. В архив входят перечисленные материалы ЛР 2 и подтверждающие материалы ЛР 1.')
    doc.add_heading('Входные данные и результаты ЛР 2',level=2)
    for relative,description in [
        ('../artifacts/manifest.csv','12 исследованных синтетических записей'),
        ('../artifacts/policy.json','политика варианта 1'),
        ('audit_findings.csv','семь находок и решения по ним'),
        ('audit_report.md','текстовый результат аудита'),
        ('data_card.md','паспорт набора данных'),
        ('model_card.md','карточка процедурного генератора'),
        ('../artifacts/passport_lr01.json','входные сведения для карточки генератора'),
        ('../artifacts/risk_register.csv','шесть рисков и меры контроля'),
        ('risk_ranking.md','ранжирование и RESULT: PASS'),
        ('validation.txt','отрицательная и положительная проверки реестра'),
        ('lab_journal.txt','ожидания до запуска и журнал выполнения'),
        ('../README.md','описание файлов и команды запуска'),
        ('../src/lab02_make_variant.py','локальный генератор учебного манифеста'),
        ('../src/lab02_audit.py','проверка лицензий и происхождения'),
        ('../src/lab02_model_card.py','сборка карточки генератора'),
        ('../src/lab02_risks.py','валидация и ранжирование рисков'),
    ]: add_link(doc,relative,description)
    doc.add_heading('Результаты эксперимента ЛР 1',level=2)
    for relative,description in [
        ('../../LR01-01/artifacts/results.json','параметры, метрики и контрольные суммы варианта 1'),
        ('../../LR01-01/artifacts/baseline_seed101.png','облачный фон 96×96, seed 101'),
        ('../../LR01-01/artifacts/factor_seed101.png','облачный фон 144×144, seed 101'),
        ('../../LR01-01/src/lab01_generator.py','исходный код процедурного генератора'),
        ('../../LR01-01/src/lab01_experiment.py','исходный код эксперимента'),
        ('../../LR01-01/reports/ЛР01_Отчет_и_паспорт.docx','отчёт и паспорт ЛР 1'),
    ]: add_link(doc,relative,description)
    doc.add_paragraph('Файлы PNG из ЛР 1 — результаты процедурной генерации облачного фона. Они не являются объектами item01–item12: в ЛР 2 исследовались записи учебного манифеста без реальных изображений.')
    doc.add_heading('Источники',level=2)
    add_link(doc,'../assignment/ЛР02_Карточка_модели_лицензии_и_риски_МУ.docx','методические указания, сценарий и критерии варианта 1')
    doc.add_paragraph('Названия групп NIST AI 600-1 приведены по перечню методических указаний. Исходные данные карточки модели подтверждаются отчётом и результатами ЛР 1, вариант 1, по ссылкам выше.')

if __name__=='__main__': main()
